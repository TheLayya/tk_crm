"""Regression checks for validated table queries before pagination."""
import json
import csv
import io
from datetime import datetime, timedelta

import pytest
from fastapi import HTTPException

from app.models.monitor import MonitorAccount, MonitorHistory, Project, ProjectMember
from app.models.op_account import OpAccount, OpAuditLog, OpCollectTask
from app.models.video import Video, VideoStats, OpAccountVideo
from app.models.team import RolePermission, UserRole
from app.models.device import Device
from app.models.proxy_node import ProxyNode
from app.services import video_service
from app.services.encryption_service import EncryptedType
from app.services.table_query_service import apply_table_rows
from tests.conftest import auth_headers


@pytest.fixture(autouse=True)
def clear_accounts(db, clean_tables):
    def clean():
        db.rollback()
        for model in (VideoStats, Video, OpAccountVideo, OpAuditLog, OpCollectTask, OpAccount,
                      MonitorHistory, MonitorAccount, ProjectMember, Project):
            db.query(model).delete()
        db.commit()
    clean()
    yield
    clean()


def test_monitor_combined_query_counts_then_pages_and_rejects_invalid_fields(client, db, super_admin):
    project = Project(name="query-monitor", created_by=super_admin.username)
    db.add(project)
    db.flush()
    for name, followers in (("Alpha%", 10), ("alpha", 20), ("Beta", 90), ("ALPHA2", 30)):
        db.add(MonitorAccount(project_id=project.id, username=name, follower_count=followers))
    db.commit()
    params = {"sort_by": "follower_count", "sort_order": "desc", "skip": 1, "limit": 1,
              "table_filters": json.dumps({"username": {"type": "text", "value": "alpha"},
                                             "follower_count": {"type": "number", "min": 15}})}
    response = client.get("/api/accounts", params=params, headers=auth_headers(super_admin))
    assert response.status_code == 200
    assert response.json()["total"] == 2
    assert [row["username"] for row in response.json()["items"]] == ["alpha"]
    params["table_filters"] = json.dumps({"username": {"type": "text", "value": "%"}})
    params["skip"] = 0
    assert client.get("/api/accounts", params=params, headers=auth_headers(super_admin)).json()["total"] == 1
    for invalid in ({"sort_by": "password"}, {"sort_order": "anything"}, {"table_filters": "[]"},
                    {"table_filters": json.dumps({"username": {"type": [], "value": "x"}})},
                    {"table_filters": json.dumps({"last_checked_at": {"type": "date", "max": "9999-12-31"}})},
                    {"table_filters": json.dumps({"last_checked_at": {"type": "date", "min": "0001-01-01T00:00:00+14:00"}})},
                    {"table_filters": json.dumps({"last_checked_at": {"type": "date", "max": "9999-12-31T23:59:59-14:00"}})},
                    {"table_filters": json.dumps({"follower_count": {"type": "number", "min": 30, "max": 1}})},
                    {"table_filters": json.dumps({"username": {"type": "number", "value": 3}})}):
        assert client.get("/api/accounts", params=invalid, headers=auth_headers(super_admin)).status_code == 422
    assert client.get("/api/accounts", params={"table_filters": json.dumps({"is_active": {"type": "enum", "value": ["yes"]}})}, headers=auth_headers(super_admin)).status_code == 422


def test_monitor_derived_query_scopes_all_rows_before_pagination(client, db, normal_user):
    role_id = db.query(UserRole.role_id).filter(UserRole.user_id == normal_user.id).scalar()
    db.add(RolePermission(role_id=role_id, permission="monitor:view"))
    owned = Project(name="query-owned", created_by=normal_user.username)
    foreign = Project(name="query-foreign", created_by="foreign")
    db.add_all([owned, foreign])
    db.flush()
    for project, name, delta in ((owned, "first", 2), (owned, "second", 8), (owned, "third", 5), (foreign, "hidden", 100)):
        account = MonitorAccount(project_id=project.id, username=name)
        db.add(account)
        db.flush()
        db.add_all([MonitorHistory(account_id=account.id, follower_count=10, check_status="success", checked_at=datetime(2026, 1, 1)),
                    MonitorHistory(account_id=account.id, follower_count=10 + delta, check_status="success", checked_at=datetime(2026, 1, 2))])
    db.commit()
    response = client.get("/api/accounts", headers=auth_headers(normal_user), params={
        "sort_by": "followers_change", "sort_order": "desc", "skip": 1, "limit": 1,
        "table_filters": json.dumps({"followers_change": {"type": "number", "min": 4},
                                      "project_name": {"type": "text", "value": "owned"}})})
    assert response.status_code == 200
    assert response.json()["total"] == 2
    assert response.json()["items"][0]["username"] == "third"
    assert response.json()["items"][0]["followers_change"] == 5


@pytest.mark.parametrize("path,boolean_field,derived_field", [
    ("/api/accounts", "is_active", "project_name"),
    ("/api/op-accounts", "tiktok_mid_video", "device_name"),
])
def test_boolean_filters_keep_validation_with_derived_queries_and_empty_results(
    client, super_admin, path, boolean_field, derived_field,
):
    for invalid in ("yes", 1):
        for extra in ({"sort_by": "followers_change"}, {}):
            filters = {boolean_field: {"type": "enum", "value": [invalid]}}
            if not extra:
                filters[derived_field] = {"type": "text", "value": "missing"}
            response = client.get(path, headers=auth_headers(super_admin), params={
                **extra, "table_filters": json.dumps(filters)})
            assert response.status_code == 422, response.text
    valid = client.get(path, headers=auth_headers(super_admin), params={
        "sort_by": "followers_change", "table_filters": json.dumps({boolean_field: {"type": "enum", "value": [True, False, None]}})})
    assert valid.status_code == 200 and valid.json() == {"items": [], "total": 0}


def test_op_account_nulls_stable_sort_and_derived_filters(client, db, normal_user):
    role_id = db.query(UserRole.role_id).filter(UserRole.user_id == normal_user.id).scalar()
    db.add(RolePermission(role_id=role_id, permission="op_account:view"))
    for name, current, previous, registrant in (("equal-first", 10, 2, normal_user.username),
                                                ("equal-second", 10, 4, normal_user.username),
                                                ("small", 2, 1, normal_user.username),
                                                ("unknown", None, None, normal_user.username),
                                                ("foreign", 100, 0, "elsewhere")):
        db.add(OpAccount(account=name, platform="tiktok", registrant=registrant,
                         follower_count=current, previous_follower_count=previous, sellers='["seller"]'))
    db.commit()
    headers = auth_headers(normal_user)
    first = client.get("/api/op-accounts", headers=headers, params={"sort_by": "follower_count", "sort_order": "desc", "limit": 1}).json()
    second = client.get("/api/op-accounts", headers=headers, params={"sort_by": "follower_count", "sort_order": "desc", "limit": 1, "skip": 1}).json()
    assert first["total"] == second["total"] == 4
    assert [first["items"][0]["account"], second["items"][0]["account"]] == ["equal-first", "equal-second"]
    null_last = client.get("/api/op-accounts", headers=headers, params={"sort_by": "follower_count", "limit": 1, "skip": 3}).json()
    assert null_last["items"][0]["account"] == "unknown"
    response = client.get("/api/op-accounts", headers=headers, params={"sort_by": "followers_change", "sort_order": "desc", "skip": 1, "limit": 1,
        "table_filters": json.dumps({"followers_change": {"type": "number", "min": 3}, "sellers": {"type": "enum", "value": ["seller"]}})})
    assert response.status_code == 200
    assert response.json()["total"] == 2
    assert response.json()["items"][0]["account"] == "equal-second"
    assert client.get("/api/op-accounts", headers=headers, params={"sort_by": "totp_secret"}).status_code == 422


def test_op_account_filter_options_cover_all_scoped_rows_without_credentials(client, db, normal_user, other_user, super_admin, monkeypatch):
    headers = auth_headers(normal_user)
    assert client.get("/api/op-accounts/filter-options").status_code == 401
    assert client.get("/api/op-accounts/filter-options", headers=headers).status_code == 403
    role_id = db.query(UserRole.role_id).filter(UserRole.user_id == normal_user.id).scalar()
    db.add(RolePermission(role_id=role_id, permission="op_account:view"))
    own_node = ProxyNode(ip="10.1.0.1", port=1080)
    foreign_node = ProxyNode(ip="10.1.0.2", port=1080)
    db.add_all([own_node, foreign_node])
    db.flush()
    own_device = Device(name="own-phone", device_type="phone", owner_id=normal_user.id)
    deleted_device = Device(name="deleted-phone", device_type="phone", owner_id=normal_user.id, is_deleted=True)
    foreign_device = Device(name="foreign-phone", device_type="phone", owner_id=other_user.id)
    db.add_all([own_device, deleted_device, foreign_device])
    db.flush()
    for index in range(55):
        db.add(OpAccount(account=f"options-{index:02}", platform="tiktok", registrant=normal_user.username,
                         country="US", purchase_channel=f"channel-{index:02}", password="secret"))
    db.add_all([
        OpAccount(account="terminal-owned", platform="tiktok", device_id=own_device.id, node_id=own_node.id,
                  country="JP", purchase_channel="last-channel", sale_customer="buyer", operator="historical-name"),
        OpAccount(account="own-blank", platform="tiktok", registrant=normal_user.username, country="", operator="",
                  purchase_channel="", sale_customer="", device_id=deleted_device.id),
        OpAccount(account="own-null", platform="tiktok", operator=normal_user.username),
        OpAccount(account="own-mail@gmail.com", platform="gmail", registrant=normal_user.username, country="MAIL"),
        OpAccount(account="foreign", platform="tiktok", registrant=other_user.username,
                  country="hidden-country", purchase_channel="hidden-channel", sale_customer="hidden-buyer",
                  device_id=foreign_device.id, node_id=foreign_node.id),
        OpAccount(account="deleted-only", platform="tiktok", device_id=deleted_device.id, country="hidden-deleted"),
    ])
    db.commit()

    def credentials_must_not_load(*args):
        raise AssertionError("filter options must not deserialize credentials")
    monkeypatch.setattr(EncryptedType, "process_result_value", credentials_must_not_load)
    response = client.get("/api/op-accounts/filter-options", headers=headers, params={"exclude_gmail": True})
    assert response.status_code == 200, response.text
    options = response.json()
    assert set(options) == {"country", "registrant", "operator", "purchase_channel", "sale_customer", "device_name", "node_ip"}
    values = lambda field: {option["value"] for option in options[field]}
    assert values("country") == {"US", "JP", "", None}
    assert values("purchase_channel") == {f"channel-{index:02}" for index in range(55)} | {"last-channel", "", None}
    assert values("device_name") == {"own-phone", None}
    assert values("node_ip") == {"10.1.0.1:1080", None}
    assert values("operator") == {"historical-name", normal_user.username, "", None}
    assert {item["value"]: item["label"] for item in options["country"]}[""] == "未填写（空白）"
    assert {item["value"]: item["label"] for item in options["device_name"]}[None] == "未绑定"
    assert "secret" not in response.text and "hidden-" not in response.text and "deleted-phone" not in response.text
    gmail = client.get("/api/op-accounts/filter-options", headers=headers).json()
    assert "MAIL" in {option["value"] for option in gmail["country"]}
    unrestricted = client.get("/api/op-accounts/filter-options", headers=auth_headers(super_admin)).json()
    assert "hidden-country" in {option["value"] for option in unrestricted["country"]}
    assert "foreign-phone" in {option["value"] for option in unrestricted["device_name"]}


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_op_account_export_matches_scoped_derived_list_filters_and_order(client, db, normal_user, format):
    role_id = db.query(UserRole.role_id).filter(UserRole.user_id == normal_user.id).scalar()
    db.add_all(RolePermission(role_id=role_id, permission=permission)
               for permission in ("op_account:view", "op_account:export", "monitor:view"))
    project = Project(name="export-project", created_by=normal_user.username)
    device = Device(name="export-phone", device_type="phone", owner_id=normal_user.id)
    node = ProxyNode(ip="10.2.0.1", port=1080)
    db.add_all([project, device, node])
    db.flush()
    start, _ = video_service._yesterday_beijing_bounds()
    for name, delta, plays in (("first", 5, 50), ("second", 9, 90)):
        account = OpAccount(account=f"export-{name}", platform="tiktok", registrant=normal_user.username,
                            device_id=device.id, node_id=node.id, password=f"secret-{name}",
                            sellers='["seller", "other"]')
        monitor = MonitorAccount(username=account.account, project_id=project.id)
        db.add_all([account, monitor])
        db.flush()
        db.add_all([MonitorHistory(account_id=monitor.id, follower_count=10, check_status="success", checked_at=start),
                    MonitorHistory(account_id=monitor.id, follower_count=10 + delta, check_status="success", checked_at=start + timedelta(hours=1)),
                    Video(account_id=monitor.id, video_id=f"export-video-{name}", play_count=plays, published_at=start + timedelta(hours=2))])
    db.add_all([
        OpAccount(account="export-unsold", platform="tiktok", registrant=normal_user.username, device_id=device.id,
                  sellers='["different"]', follower_count=100, previous_follower_count=0),
        OpAccount(account="foreign-export", platform="tiktok", registrant="foreign", sellers='["seller"]',
                  follower_count=1000, previous_follower_count=0),
    ])
    db.commit()
    headers = auth_headers(normal_user)
    params = {"sort_by": "followers_change", "sort_order": "desc", "exclude_gmail": True,
              "table_filters": json.dumps({"platform": {"type": "enum", "value": ["tiktok"]},
                                              "sellers": {"type": "enum", "value": ["seller"]},
                                              "device_name": {"type": "enum", "value": [device.name]},
                                              "node_ip": {"type": "enum", "value": ["10.2.0.1:1080"]},
                                              "followers_change": {"type": "number", "min": 4},
                                              "video_source": {"type": "enum", "value": ["monitor"]},
                                              "yesterday_video_count": {"type": "number", "min": 1},
                                              "yesterday_video_plays": {"type": "number", "min": 40}})}
    expected = client.get("/api/op-accounts", headers=headers, params=params)
    assert expected.status_code == 200, expected.text
    expected_names = [row["account"] for row in expected.json()["items"]]
    assert expected_names == ["export-second", "export-first"] and expected.json()["total"] == 2
    page = client.get("/api/op-accounts", headers=headers, params={**params, "skip": 1, "limit": 1}).json()
    assert page["total"] == 2 and [row["account"] for row in page["items"]] == expected_names[1:]
    exported = client.get("/api/op-accounts/export", headers=headers, params={**params, "format": format})
    assert exported.status_code == 200, exported.text
    if format == "csv":
        rows = list(csv.DictReader(io.StringIO(exported.content.decode("utf-8-sig"))))
    else:
        import openpyxl
        sheet = openpyxl.load_workbook(io.BytesIO(exported.content), read_only=True).active
        values = list(sheet.values)
        rows = [dict(zip(values[0], row)) for row in values[1:]]
    assert [row["account"] for row in rows] == expected_names
    assert [row["password"] for row in rows] == ["secret-second", "secret-first"]


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_op_account_export_validates_queries_when_empty(client, super_admin, format):
    headers = auth_headers(super_admin)
    for invalid in ({"sort_by": "password"}, {"sort_by": "missing"}, {"sort_order": "invalid"},
                    {"table_filters": "[]"}, {"table_filters": '{"tiktok_mid_video":{"type":"enum","value":["yes"]}}'},
                    {"sort_by": "followers_change", "table_filters": '{"follower_count":{"type":"number","min":"NaN"}}'}):
        response = client.get("/api/op-accounts/export", headers=headers, params={"format": format, **invalid})
        assert response.status_code == 422, response.text
    params = {"sort_by": "followers_change", "table_filters": '{"sellers":{"type":"enum","value":["missing"]}}'}
    assert client.get("/api/op-accounts", headers=headers, params=params).json() == {"items": [], "total": 0}
    response = client.get("/api/op-accounts/export", headers=headers, params={"format": format, **params})
    assert response.status_code == 200
    if format == "csv":
        assert list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig")))) == []
    else:
        import openpyxl
        assert openpyxl.load_workbook(io.BytesIO(response.content), read_only=True).active.max_row == 1


def test_video_and_history_queries_use_filtered_totals_and_date_ranges(client, db):
    project = Project(name="query-video")
    db.add(project)
    db.flush()
    account = MonitorAccount(project_id=project.id, username="query-video")
    db.add(account)
    db.flush()
    for index in range(3):
        when = datetime(2026, 1, 1, 12) + timedelta(days=index)
        db.add(Video(account_id=account.id, video_id=f"query-{index}", play_count=10 + index, published_at=when))
        db.add(MonitorHistory(account_id=account.id, check_status="success", follower_count=10 + index, checked_at=when))
    db.commit()
    params = {"sort_by": "play_count", "sort_order": "desc", "skip": 1, "limit": 1,
              "table_filters": json.dumps({"play_count": {"type": "number", "min": 11},
                                             "published_at": {"type": "date", "max": "2026-01-03"}})}
    videos = client.get(f"/api/accounts/{account.id}/videos", params=params).json()
    assert videos["total"] == 2
    assert videos["items"][0]["play_count"] == 11
    history = client.get(f"/api/accounts/{account.id}/history", params={"sort_by": "follower_count", "sort_order": "desc",
        "table_filters": json.dumps({"checked_at": {"type": "date", "value": "2026-01-02"}})})
    assert history.status_code == 200 and [row["follower_count"] for row in history.json()] == [11]
    assert client.get(f"/api/accounts/{account.id}/history", params={"sort_by": "missing"}).status_code == 422


def test_complete_rows_match_enums_and_keep_null_last_on_both_sort_directions():
    fields = {"id": "number", "amount": "number", "members": "text", "time": "date"}
    rows = [{"id": 3, "amount": None, "members": ["Amy"], "time": None},
            {"id": 2, "amount": 1, "members": ["Amy"], "time": datetime(2026, 1, 1, 20)},
            {"id": 1, "amount": 1, "members": ["Ben"], "time": datetime(2026, 1, 1, 12)}]
    for direction in ("asc", "desc"):
        result = apply_table_rows(rows, fields, "amount", direction)
        assert [row["id"] for row in result] == [1, 2, 3]
    result = apply_table_rows(rows, fields, table_filters={"members": {"type": "enum", "value": ["Amy"]},
                                                         "time": {"type": "date", "max": "2026-01-01"}})
    assert [row["id"] for row in result] == [2]
    with pytest.raises(HTTPException) as exc:
        apply_table_rows([], fields, table_filters={"amount": {"type": "number", "value": "NaN"}})
    assert exc.value.status_code == 422


def test_numeric_video_arrays_sort_by_sum_and_work_item_queries(client, db, super_admin):
    from app.models.work_item import WorkItem
    fields = {"id": "number", "plays": "number"}
    rows = [{"id": 1, "plays": [20]}, {"id": 2, "plays": [100, 20]}, {"id": 3, "plays": None}]
    assert [row["id"] for row in apply_table_rows(rows, fields, "plays", "desc")] == [2, 1, 3]
    assert [row["id"] for row in apply_table_rows(rows, fields, table_filters={"plays": {"type": "number", "min": 50}})] == [2]
    db.add_all([WorkItem(title="Zulu", category="其他", is_done=False, created_by=super_admin.username, reminder_users=",root,"),
                WorkItem(title="Alpha", category="其他", is_done=False, created_by=super_admin.username, reminder_users=",elsewhere,")])
    db.commit()
    result = client.get("/api/work-items", headers=auth_headers(super_admin), params={
        "sort_by": "title", "sort_order": "asc", "table_filters": json.dumps({"reminder_users": {"type": "text", "value": "root"}})})
    assert result.status_code == 200 and result.json()["total"] == 1
    assert result.json()["items"][0]["title"] == "Zulu"
    invalid = client.get("/api/work-items", headers=auth_headers(super_admin), params={
        "table_filters": json.dumps({"is_done": {"type": "enum", "value": ["yes"]}})})
    assert invalid.status_code == 422
