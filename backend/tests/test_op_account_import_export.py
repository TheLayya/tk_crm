import io

import openpyxl

from app.models.op_account import OpAccount
from app.services.op_account_service import create_import_template, export_op_accounts, import_from_excel


def test_excel_round_trip_preserves_editable_fields(db):
    account = OpAccount(
        platform="tiktok",
        account="roundtrip-source",
        password="secret",
        status="正常",
        tiktok_mid_video=True,
        tiktok_showcase=False,
    )
    db.add(account)
    db.commit()

    exported = export_op_accounts(db, {}, format="xlsx")
    workbook = openpyxl.load_workbook(io.BytesIO(exported))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    values = [cell.value for cell in sheet[2]]
    values[headers.index("account")] = "roundtrip-copy"
    values[headers.index("tiktok_mid_video")] = "0"
    values[headers.index("tiktok_showcase")] = "1"
    sheet.append(values)
    output = io.BytesIO()
    workbook.save(output)

    result = import_from_excel(db, output.getvalue())
    assert result.success == 1
    copy = db.query(OpAccount).filter(OpAccount.account == "roundtrip-copy").one()
    assert copy.tiktok_mid_video is False
    assert copy.tiktok_showcase is True


def test_excel_import_rejects_missing_required_headers(db):
    workbook = openpyxl.Workbook()
    workbook.active.append(["nickname"])
    output = io.BytesIO()
    workbook.save(output)

    try:
        import_from_excel(db, output.getvalue())
    except ValueError as exc:
        assert "platform" in str(exc)
    else:
        raise AssertionError("missing required headers should be rejected")


def test_chinese_excel_template_can_be_imported(db):
    template = create_import_template()
    workbook = openpyxl.load_workbook(io.BytesIO(template))
    sheet = workbook["运营账号导入"]
    headers = [cell.value for cell in sheet[1]]
    values = [cell.value for cell in sheet[2]]
    values[headers.index("账号")] = "chinese-template-account"
    values[headers.index("平台")] = "TikTok"
    sheet.delete_rows(2)
    sheet.append(values)
    output = io.BytesIO()
    workbook.save(output)

    result = import_from_excel(db, output.getvalue())
    assert result.success == 1
    account = db.query(OpAccount).filter(OpAccount.account == "chinese-template-account").one()
    assert account.platform == "tiktok"
    assert account.tiktok_mid_video is False


def test_localized_export_uses_chinese_headers(db):
    db.add(OpAccount(platform="tiktok", account="localized-export", status="正常"))
    db.commit()

    exported = export_op_accounts(db, {}, format="xlsx", localized=True)
    workbook = openpyxl.load_workbook(io.BytesIO(exported), read_only=True)
    headers = [cell.value for cell in workbook.active[1]]
    assert headers[:2] == ["账号", "平台"]
