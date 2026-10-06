"""Isolated API regressions for the monitor administration workflows."""
import csv
import io
import threading
from datetime import datetime, timedelta
from unittest.mock import AsyncMock

import openpyxl
import pytest

from app.models.monitor import (
    MonitorAccount, MonitorHistory, MonitorProxy, MonitorSettings, Project, ProjectMember,
)
from app.models.team import Department, Role, RolePermission, UserRole
from app.models.video import Video, VideoStats
from app.services import monitor_service, proxy_service, video_service
from tests.conftest import auth_headers


@pytest.fixture(autouse=True)
def isolated_monitor_tables(db, clean_tables, monkeypatch):
    def clean():
        db.rollback()
        for model in (VideoStats, Video, MonitorHistory, MonitorAccount, ProjectMember,
                      Project, MonitorProxy, MonitorSettings):
            db.query(model).delete()
        db.commit()

    clean()
    original_start = threading.Thread.start

    def start(thread):
        if getattr(thread._target, "__name__", "") != "_trigger_first_check":
            original_start(thread)

    monkeypatch.setattr(threading.Thread, "start", start)
    yield
    clean()


def make_project(db, name="monitor", owner="root"):
    project = Project(name=name, created_by=owner)
    db.add(project)
    db.commit()
    return project


def grant_monitor(db, user, scope="self"):
    role = Role(name=f"monitor-{user.username}-{scope}", data_scope=scope)
    db.add(role)
    db.flush()
    db.add_all([RolePermission(role_id=role.id, permission="monitor:view"),
                UserRole(user_id=user.id, role_id=role.id)])
    db.commit()


def test_project_crud_conflicts_members_and_account_count(client, db, super_admin):
    headers = auth_headers(super_admin)
    created = client.post("/api/projects", headers=headers, json={"name": "项目", "description": "initial"})
    assert created.status_code == 201, created.text
    project_id = created.json()["id"]
    assert created.json()["created_by"] == super_admin.username
    assert created.json()["account_count"] == 0
    assert client.post("/api/projects", headers=headers, json={"name": "项目"}).status_code == 409
    second = client.post("/api/projects", headers=headers, json={"name": "second"}).json()
    assert client.put(f"/api/projects/{project_id}", headers=headers, json={"name": "second"}).status_code == 409
    updated = client.put(f"/api/projects/{project_id}", headers=headers, json={"description": "updated"})
    assert updated.status_code == 200 and updated.json()["description"] == "updated"
    members = client.put(f"/api/projects/{project_id}/members", headers=headers, json={"usernames": ["alice", "alice", "bob"]})
    assert members.status_code == 200 and set(members.json()) == {"alice", "bob"}
    assert set(client.get(f"/api/projects/{project_id}/members", headers=headers).json()) == {"alice", "bob"}
    account = MonitorAccount(project_id=project_id, username="included")
    db.add(account)
    db.commit()
    assert client.get(f"/api/projects/{project_id}", headers=headers).json()["account_count"] == 1
    listing = client.get("/api/projects", headers=headers)
    assert listing.status_code == 200 and len(listing.json()) == 2
    assert client.delete(f"/api/projects/{project_id}", headers=headers).status_code == 409
    assert client.delete(f"/api/projects/{second['id']}", headers=headers).status_code == 204
    assert db.query(Project).filter(Project.id == second["id"]).first() is None
    for method, suffix, payload in [("get", "", None), ("put", "", {"name": "missing"}),
                                    ("delete", "", None), ("get", "/members", None),
                                    ("put", "/members", {"usernames": []})]:
        response = client.request(method, "/api/projects/999999" + suffix, headers=headers, json=payload)
        assert response.status_code == 404, response.text


def test_member_can_read_invited_project_but_only_edit_owned_project(client, db, normal_user, other_user):
    grant_monitor(db, normal_user)
    own = make_project(db, "own", normal_user.username)
    invited = make_project(db, "invited", other_user.username)
    hidden = make_project(db, "hidden", other_user.username)
    db.add(ProjectMember(project_id=invited.id, username=normal_user.username))
    db.commit()
    headers = auth_headers(normal_user)
    listing = client.get("/api/projects", headers=headers)
    assert listing.status_code == 200 and {row["id"] for row in listing.json()} == {own.id, invited.id}
    assert client.get(f"/api/projects/{invited.id}", headers=headers).status_code == 200
    assert client.get(f"/api/projects/{hidden.id}", headers=headers).status_code == 403
    assert client.get(f"/api/projects/{hidden.id}/members", headers=headers).status_code == 403
    for method, suffix, payload in [("put", "", {"description": "forbidden"}),
                                    ("delete", "", None), ("put", "/members", {"usernames": []})]:
        assert client.request(method, f"/api/projects/{invited.id}" + suffix, headers=headers, json=payload).status_code == 403
    assert client.put(f"/api/projects/{own.id}", headers=headers, json={"name": "renamed"}).status_code == 200
    db.refresh(invited)
    assert invited.description is None
    assert db.query(ProjectMember).count() == 1


def test_department_monitor_scope_only_returns_department_projects(client, db, normal_user, other_user):
    department = Department(name="monitor-dept")
    db.add(department)
    db.flush()
    normal_user.department_id = department.id
    other_user.department_id = department.id
    db.commit()
    grant_monitor(db, normal_user, "dept")
    own = make_project(db, "own-dept", normal_user.username)
    colleague = make_project(db, "colleague", other_user.username)
    make_project(db, "external", "external")
    response = client.get("/api/projects", headers=auth_headers(normal_user))
    assert response.status_code == 200
    assert {row["id"] for row in response.json()} == {own.id, colleague.id}


def test_monitor_account_crud_filters_and_delete_cascade(client, db, super_admin):
    project = make_project(db)
    headers = auth_headers(super_admin)
    body = {"project_id": project.id, "username": "first", "use_proxy": False}
    response = client.post("/api/accounts", json=body)
    assert response.status_code == 201, response.text
    account_id = response.json()["id"]
    assert response.json()["monitor_interval"] == 3600 and response.json()["project_name"] == project.name
    assert client.post("/api/accounts", json=body).status_code == 409
    assert client.post("/api/accounts", json={**body, "username": "missing-proxy", "proxy_id": 99999}).status_code == 404
    assert client.put(f"/api/accounts/{account_id}", json={"proxy_id": 99999}).status_code == 404
    updated = client.put(f"/api/accounts/{account_id}", json={"nickname": "Match nickname", "is_active": False, "monitor_interval": 600})
    assert updated.status_code == 200 and updated.json()["nickname"] == "Match nickname"
    assert updated.json()["monitor_interval"] == 600
    assert client.get(f"/api/accounts/{account_id}").json()["username"] == "first"
    db.add_all([MonitorAccount(project_id=project.id, username="second"),
                MonitorHistory(account_id=account_id, check_status="failed", error_message="ACCOUNT_NOT_FOUND:gone"),
                Video(account_id=account_id, video_id="cascade")])
    db.commit()
    video = db.query(Video).one()
    db.add(VideoStats(video_id=video.id, play_count=20))
    db.commit()
    listing = client.get("/api/accounts", headers=headers, params={"keyword": "Match", "is_active": False, "project_id": project.id})
    assert listing.status_code == 200 and listing.json()["total"] == 1
    assert listing.json()["items"][0]["latest_check_status"] == "not_found"
    page = client.get("/api/accounts", headers=headers, params={"skip": 1, "limit": 1})
    assert page.status_code == 200 and page.json()["total"] == 2
    assert len(page.json()["items"]) == 1
    assert client.delete(f"/api/accounts/{account_id}").status_code == 204
    assert db.query(MonitorHistory).count() == db.query(Video).count() == db.query(VideoStats).count() == 0
    for method, payload in [("get", None), ("put", {"nickname": "missing"}), ("delete", None), ("post", None)]:
        suffix = "/check" if method == "post" else ""
        assert client.request(method, "/api/accounts/999999" + suffix, json=payload).status_code == 404


def test_monitor_list_hides_foreign_project_and_includes_invited_accounts(client, db, normal_user, other_user):
    grant_monitor(db, normal_user)
    visible = make_project(db, "visible", other_user.username)
    hidden = make_project(db, "secret", other_user.username)
    db.add_all([ProjectMember(project_id=visible.id, username=normal_user.username),
                MonitorAccount(project_id=visible.id, username="visible-account"),
                MonitorAccount(project_id=hidden.id, username="hidden-account")])
    db.commit()
    headers = auth_headers(normal_user)
    result = client.get("/api/accounts", headers=headers).json()
    assert result["total"] == 1 and result["items"][0]["username"] == "visible-account"
    assert client.get("/api/accounts", headers=headers, params={"project_id": hidden.id}).json() == {"items": [], "total": 0}


def test_text_import_validates_rows_and_batch_actions_persist(client, db):
    source = make_project(db, "batch-source")
    target = make_project(db, "batch-target")
    result = client.post("/api/accounts/import", json={"project_id": source.id,
        "usernames": "valid.one\nvalid_two\nvalid.one\nbad name\n127.0.0.1:1080\n"})
    assert result.status_code == 200, result.text
    assert {key: result.json()[key] for key in ("total", "success", "duplicates", "failed")} == {
        "total": 5, "success": 2, "duplicates": 1, "failed": 2,
    }
    assert "代理" in result.json()["results"][4]["reason"]
    rows = db.query(MonitorAccount).all()
    ids = [row.id for row in rows]
    for action, expected in [("disable", False), ("enable", True)]:
        response = client.post("/api/accounts/batch", json={"account_ids": ids + [99999], "action": action})
        assert response.status_code == 200 and response.json()["updated"] == 2
        db.expire_all()
        assert {row.is_active for row in rows} == {expected}
    assert client.post("/api/accounts/batch", json={"account_ids": ids, "action": "move"}).status_code == 422
    assert client.post("/api/accounts/batch", json={"account_ids": ids, "action": "move", "target_project_id": 99999}).status_code == 404
    moved = client.post("/api/accounts/batch", json={"account_ids": ids, "action": "move", "target_project_id": target.id})
    assert moved.status_code == 200 and moved.json()["moved"] == 2
    db.expire_all()
    assert {row.project_id for row in rows} == {target.id}
    assert client.post("/api/accounts/batch", json={"account_ids": ids, "action": "invalid"}).status_code == 422
    deleted = client.post("/api/accounts/batch", json={"account_ids": ids + [99999], "action": "delete"})
    assert deleted.status_code == 200 and deleted.json()["deleted"] == 2
    assert db.query(MonitorAccount).count() == 0


def test_manual_account_check_records_profile_history(client, db, monkeypatch):
    project = make_project(db)
    account = MonitorAccount(project_id=project.id, username="manual", use_proxy=False, enable_video_monitoring=False)
    db.add(account)
    db.commit()
    fetch = AsyncMock(return_value={"success": True, "data": {"follower_count": 27, "nickname": "Manual result"}})
    monkeypatch.setattr(monitor_service.scraper_service, "fetch_user_info", fetch)
    response = client.post(f"/api/accounts/{account.id}/check")
    assert response.status_code == 202
    history = db.query(MonitorHistory).one()
    assert response.json()["history_id"] == history.id
    assert history.follower_count == 27 and history.check_status == "success"
    db.refresh(account)
    assert account.nickname == "Manual result"
    fetch.assert_awaited_once_with("manual", proxy=None)


def test_history_date_window_pagination_and_trend_deltas(client, db):
    project = make_project(db)
    account = MonitorAccount(project_id=project.id, username="trend")
    db.add(account)
    db.flush()
    start = datetime(2026, 1, 1)
    for day, followers, likes in [(0, 10, 20), (1, 15, 40), (2, 12, 45)]:
        db.add(MonitorHistory(account_id=account.id, follower_count=followers, like_count=likes,
                              checked_at=start + timedelta(days=day), check_status="success"))
    db.commit()
    response = client.get(f"/api/accounts/{account.id}/history", params={
        "start_time": start.isoformat(), "end_time": (start + timedelta(days=2)).isoformat(),
        "skip": 1, "limit": 1,
    })
    assert response.status_code == 200 and [row["follower_count"] for row in response.json()] == [15]
    trends = client.get(f"/api/accounts/{account.id}/trends", params={
        "start_time": start.isoformat(), "end_time": (start + timedelta(days=2)).isoformat(),
    })
    assert trends.status_code == 200
    points = trends.json()["data_points"]
    assert [point["followers_change"] for point in points] == [0, 5, -3]
    assert [point["likes_change"] for point in points] == [0, 20, 5]
    assert client.get(f"/api/accounts/{account.id}/trends", params={"start_time": "2030-01-01T00:00:00"}).json() == {"data_points": []}
    for suffix in ["history", "trends"]:
        assert client.get(f"/api/accounts/99999/{suffix}").status_code == 404


def test_proxy_crud_batch_parsing_and_mock_connectivity_counters(client, db, monkeypatch):
    response = client.post("/api/proxies", json={"host": "127.0.0.1", "port": 1080, "username": "user", "password": "pass"})
    assert response.status_code == 201
    proxy_id = response.json()["id"]
    assert response.json()["name"] == "user@127.0.0.1:1080"
    assert client.get(f"/api/proxies/{proxy_id}").json()["username"] == "user"
    changed = client.put(f"/api/proxies/{proxy_id}", json={"host": "127.0.0.2", "username": None})
    assert changed.status_code == 200 and changed.json()["name"] == "127.0.0.2:1080"
    batch = client.post("/api/proxies/batch", json={"proxies_text": "10.0.0.1:80\n10.0.0.2:1080:u:p\nbad\n10.0.0.3:wrong"})
    assert batch.status_code == 201 and (batch.json()["success_count"], batch.json()["fail_count"]) == (2, 2)
    assert len(batch.json()["errors"]) == 2
    listing = client.get("/api/proxies", params={"skip": 1, "limit": 1})
    assert listing.status_code == 200 and len(listing.json()) == 1
    fetch = AsyncMock(side_effect=[{"success": True, "response_time": 1.25},
                                  {"success": False, "response_time": None, "error": "timeout"}])
    monkeypatch.setattr(proxy_service.scraper_service, "test_proxy", fetch)
    assert client.post(f"/api/proxies/{proxy_id}/test").json() == {"success": True, "response_time": 1.25, "error": None}
    assert client.post(f"/api/proxies/{proxy_id}/test").json()["error"] == "timeout"
    saved = db.query(MonitorProxy).filter(MonitorProxy.id == proxy_id).one()
    assert (saved.success_count, saved.fail_count, saved.last_test_result) == (1, 1, "failed")
    assert saved.last_test_at is not None and fetch.await_count == 2
    assert client.delete(f"/api/proxies/{proxy_id}").status_code == 204
    assert db.query(MonitorProxy).filter(MonitorProxy.id == proxy_id).first() is None
    for method, suffix, payload in [("get", "", None), ("put", "", {"host": "missing"}),
                                    ("delete", "", None), ("post", "/test", None)]:
        assert client.request(method, "/api/proxies/99999" + suffix, json=payload).status_code == 404


def test_video_api_paginates_snapshots_and_manual_collection(client, db, monkeypatch):
    project = make_project(db)
    account = MonitorAccount(project_id=project.id, username="videos", sec_uid="sec", use_proxy=False)
    db.add(account)
    db.flush()
    old = Video(account_id=account.id, video_id="old", published_at=datetime(2020, 1, 1), play_count=10)
    new = Video(account_id=account.id, video_id="new", published_at=datetime(2020, 1, 2), play_count=20)
    db.add_all([old, new])
    db.flush()
    db.add_all([VideoStats(video_id=new.id, play_count=5, recorded_at=datetime(2020, 1, 1)),
                VideoStats(video_id=new.id, play_count=20, recorded_at=datetime(2020, 1, 2))])
    db.commit()
    response = client.get(f"/api/accounts/{account.id}/videos", params={"skip": 1, "limit": 1})
    assert response.status_code == 200 and response.json()["total"] == 2
    assert [item["video_id"] for item in response.json()["items"]] == ["old"]
    assert [row["play_count"] for row in client.get(f"/api/videos/{new.id}/stats").json()] == [5, 20]
    fetch = AsyncMock(return_value={"success": True, "data": [{"video_id": "manual-new", "play_count": 30}]})
    monkeypatch.setattr(video_service.scraper_service, "fetch_user_videos", fetch)
    collected = client.post(f"/api/accounts/{account.id}/videos/collect")
    assert collected.status_code == 202 and collected.json()["new_videos"] == 1
    assert db.query(Video).filter(Video.video_id == "manual-new").one().play_count == 30
    assert db.query(VideoStats).count() == 3
    fetch.assert_awaited_once_with("sec", proxy=None)
    assert client.get("/api/accounts/99999/videos").status_code == 404
    assert client.post("/api/accounts/99999/videos/collect").status_code == 404
    assert client.get("/api/videos/99999/stats").status_code == 404


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_file_import_and_export_preserve_account_rows(client, db, extension):
    project = make_project(db, "导出项目")
    if extension == "csv":
        payload = "username,other\n@first,1\nsecond,2\nfirst,3\n".encode("utf-8-sig")
    else:
        workbook = openpyxl.Workbook()
        for row in [["UserName", "other"], ["@first", 1], ["second", 2], ["first", 3]]:
            workbook.active.append(row)
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        payload = output.getvalue()
    response = client.post("/api/accounts/import-file", params={"project_id": project.id, "monitor_interval": 600},
                           files={"file": (f"accounts.{extension}", payload)})
    assert response.status_code == 200, response.text
    assert (response.json()["success"], response.json()["duplicates"]) == (2, 1)
    assert {row.monitor_interval for row in db.query(MonitorAccount).all()} == {600}
    exported = client.get("/api/accounts/export", params={"project_id": project.id, "format": "csv"})
    assert exported.status_code == 200 and exported.content.startswith(b"\xef\xbb\xbf")
    assert "filename*=UTF-8''" in exported.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(exported.content.decode("utf-8-sig"))))
    assert [row["username"] for row in rows] == ["first", "second"]
    excel = client.get("/api/accounts/export", params={"format": "excel"})
    assert excel.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(excel.content), data_only=True)
    try:
        assert workbook.active.max_row == 3
        assert workbook.active.cell(2, 1).value == "first"
    finally:
        workbook.close()


def test_file_import_rejects_missing_project_empty_and_invalid_formats(client, db):
    project = make_project(db)
    for project_id, filename, payload, status in [(99999, "accounts.csv", b"username\nfirst", 404),
                                                (project.id, "accounts.csv", b"", 422),
                                                (project.id, "accounts.txt", b"first", 422),
                                                (project.id, "accounts.csv", b"wrong\nfirst", 422)]:
        response = client.post("/api/accounts/import-file", params={"project_id": project_id}, files={"file": (filename, payload)})
        assert response.status_code == status, response.text
    assert db.query(MonitorAccount).count() == 0
    assert client.get("/api/accounts/export", params={"project_id": 99999}).status_code == 404
    assert client.get("/api/accounts/export", params={"format": "invalid"}).status_code == 422
